from __future__ import annotations

import platform
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from config.constants import CUT_SEW_WORKBOOK_PATH


CELL_ADDRESS_PATTERN = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]*$")


def resolve_cut_sew_workbook_path(workbook_path: str | None = None) -> Path:
    return Path(workbook_path or CUT_SEW_WORKBOOK_PATH)


def cut_sew_workbook_exists(workbook_path: str | Path | None = None) -> bool:
    selected_path = resolve_cut_sew_workbook_path(str(workbook_path) if workbook_path else None)
    return selected_path.exists()


def _visible_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _has_visible_value(row_values: list[object]) -> bool:
    return any(_visible_value(value) for value in row_values)


def _trim_empty_edges(rows: list[list[object]]) -> list[list[object]]:
    if not rows:
        return rows

    max_column_index = 0
    for row_values in rows:
        for column_index, value in enumerate(row_values, start=1):
            if _visible_value(value):
                max_column_index = max(max_column_index, column_index)

    if max_column_index == 0:
        return []

    return [row_values[:max_column_index] for row_values in rows]


def _build_unique_headers(header_values: list[object]) -> list[str]:
    headers: list[str] = []
    used_headers: dict[str, int] = {}

    for index, value in enumerate(header_values, start=1):
        base_header = str(value).strip() if _visible_value(value) else f"Column_{index}"
        duplicate_count = used_headers.get(base_header, 0)
        used_headers[base_header] = duplicate_count + 1
        header_name = base_header if duplicate_count == 0 else f"{base_header}_{duplicate_count + 1}"
        headers.append(header_name)

    return headers


def _collect_non_empty_rows(worksheet) -> list[tuple[int, list[object]]]:
    non_empty_rows: list[tuple[int, list[object]]] = []
    for row in worksheet.iter_rows(values_only=False):
        row_values = [cell.value for cell in row]
        if _has_visible_value(row_values):
            non_empty_rows.append((row[0].row, row_values))
    return non_empty_rows


def _worksheet_to_dataframe(worksheet) -> tuple[pd.DataFrame, dict[str, object]]:
    row_items = _collect_non_empty_rows(worksheet)
    if not row_items:
        return pd.DataFrame(), {}

    row_numbers = [row_number for row_number, _ in row_items]
    row_values = [values for _, values in row_items]
    trimmed_rows = _trim_empty_edges(row_values)
    if not trimmed_rows:
        return pd.DataFrame(), {}

    headers = _build_unique_headers(trimmed_rows[0])
    header_excel_row = row_numbers[0]

    data_rows = []
    data_excel_rows = []
    for row_number, row_values in zip(row_numbers[1:], trimmed_rows[1:]):
        padded_row = row_values + [None] * (len(headers) - len(row_values))
        data_rows.append(padded_row[: len(headers)])
        data_excel_rows.append(row_number)

    dataframe = pd.DataFrame(data_rows, columns=headers)
    dataframe["__excel_row_number"] = data_excel_rows
    dataframe = dataframe.dropna(how="all", subset=headers)

    column_excel_indexes = {column_name: index for index, column_name in enumerate(headers, start=1)}
    metadata = {
        "header_excel_row": header_excel_row,
        "column_excel_indexes": column_excel_indexes,
    }
    return dataframe, metadata


def load_cut_sew_tables(
    workbook_path: str | Path | None = None,
    *,
    data_only: bool = True,
) -> list[dict[str, object]]:
    selected_path = resolve_cut_sew_workbook_path(str(workbook_path) if workbook_path else None)
    workbook = load_workbook(selected_path, read_only=True, data_only=data_only)
    tables: list[dict[str, object]] = []

    try:
        for worksheet in workbook.worksheets:
            dataframe, metadata = _worksheet_to_dataframe(worksheet)
            if dataframe.empty:
                continue

            visible_columns = [column for column in dataframe.columns if not str(column).startswith("__")]
            tables.append(
                {
                    "sheet_name": worksheet.title,
                    "dataframe": dataframe,
                    "display_dataframe": dataframe[visible_columns].copy(),
                    "row_count": len(dataframe),
                    "column_count": len(visible_columns),
                    "metadata": metadata,
                }
            )
    finally:
        workbook.close()

    return tables


def get_cut_sew_sheet_names(workbook_path: str | Path | None = None) -> list[str]:
    selected_path = resolve_cut_sew_workbook_path(str(workbook_path) if workbook_path else None)
    workbook = load_workbook(selected_path, read_only=True, data_only=False)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def validate_cell_address(cell_address: str) -> str:
    cleaned_address = str(cell_address).strip().upper()
    if not CELL_ADDRESS_PATTERN.match(cleaned_address):
        raise ValueError("Please enter a valid Excel cell address like D4, J14, or AA20.")
    return cleaned_address


def get_cell_value_pair(
    workbook_path: str | Path,
    sheet_name: str,
    cell_address: str,
) -> dict[str, Any]:
    selected_path = resolve_cut_sew_workbook_path(str(workbook_path))
    cleaned_address = validate_cell_address(cell_address)

    formula_workbook = load_workbook(selected_path, read_only=True, data_only=False)
    value_workbook = load_workbook(selected_path, read_only=True, data_only=True)
    try:
        if sheet_name not in formula_workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")

        formula_value = formula_workbook[sheet_name][cleaned_address].value
        calculated_value = value_workbook[sheet_name][cleaned_address].value
        return {
            "cell": cleaned_address,
            "formula_or_input": formula_value,
            "calculated_value": calculated_value,
            "is_formula": isinstance(formula_value, str) and formula_value.startswith("="),
        }
    finally:
        formula_workbook.close()
        value_workbook.close()


def _coerce_user_value(raw_value: object) -> object:
    if raw_value is None or pd.isna(raw_value):
        return None

    value = str(raw_value).strip()
    if value == "":
        return None

    normalized_value = value.replace(",", "")
    try:
        numeric_value = float(normalized_value)
        return int(numeric_value) if numeric_value.is_integer() else numeric_value
    except ValueError:
        return value


def _recalculate_with_excel_com(workbook_path: Path) -> tuple[bool, str]:
    if platform.system().lower() != "windows":
        return False, "Excel recalculation is available only on Windows with Microsoft Excel installed."

    try:
        import win32com.client  # type: ignore
    except ImportError:
        return False, "pywin32 is not installed. Run: pip install pywin32"

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(workbook_path.resolve()))
        excel.CalculateFullRebuild()
        workbook.Save()
        workbook.Close(SaveChanges=True)
        excel.Quit()
        return True, "Excel recalculation completed successfully."
    except Exception as exc:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()
        return False, f"Excel recalculation failed: {exc}"


def _find_value_column_index(worksheet, header_excel_row: int, value_column_name: str = "Value") -> int:
    for cell in worksheet[header_excel_row]:
        if str(cell.value).strip().lower() == value_column_name.lower():
            return cell.column
    raise ValueError(f"Could not find '{value_column_name}' column in sheet {worksheet.title}.")


def update_input_values_and_recalculate(
    workbook_path: str | Path,
    sheet_name: str,
    header_excel_row: int,
    row_value_updates: dict[int, object],
    *,
    value_column_name: str = "Value",
) -> dict[str, object]:
    selected_path = resolve_cut_sew_workbook_path(str(workbook_path))

    if not row_value_updates:
        return {"updated_count": 0, "recalculated": False, "message": "No input changes found."}

    workbook = load_workbook(selected_path, read_only=False, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")

        worksheet = workbook[sheet_name]
        value_column_index = _find_value_column_index(worksheet, header_excel_row, value_column_name)
        updated_count = 0

        for excel_row_number, raw_new_value in row_value_updates.items():
            target_cell = worksheet.cell(row=int(excel_row_number), column=value_column_index)
            existing_value = target_cell.value

            if isinstance(existing_value, str) and existing_value.startswith("="):
                cell_address = f"{get_column_letter(value_column_index)}{excel_row_number}"
                raise ValueError(
                    f"{sheet_name}!{cell_address} is a formula cell. Only input rows can be edited."
                )

            target_cell.value = _coerce_user_value(raw_new_value)
            updated_count += 1

        workbook.save(selected_path)
    finally:
        workbook.close()

    recalculated, message = _recalculate_with_excel_com(selected_path)
    return {
        "updated_count": updated_count,
        "recalculated": recalculated,
        "message": message,
    }


def update_cell_and_recalculate(
    workbook_path: str | Path,
    sheet_name: str,
    cell_address: str,
    raw_new_value: str,
) -> dict[str, object]:
    selected_path = resolve_cut_sew_workbook_path(str(workbook_path))
    cleaned_address = validate_cell_address(cell_address)
    new_value = _coerce_user_value(raw_new_value)

    workbook = load_workbook(selected_path, read_only=False, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")

        worksheet = workbook[sheet_name]
        current_value = worksheet[cleaned_address].value
        if isinstance(current_value, str) and current_value.startswith("="):
            raise ValueError(
                f"{sheet_name}!{cleaned_address} is a formula cell. Please edit only input/driver cells."
            )

        worksheet[cleaned_address] = new_value
        workbook.save(selected_path)
    finally:
        workbook.close()

    recalculated, message = _recalculate_with_excel_com(selected_path)
    return {
        "updated_cell": f"{sheet_name}!{cleaned_address}",
        "updated_value": new_value,
        "recalculated": recalculated,
        "message": message,
    }
